# -*- coding: utf-8 -*-
"""
自动化脚本：批量读取 Excel 客户信息，按意向等级生成个性化跟进消息
=================================================================
用法：python 批量跟进消息生成.py [输入.xlsx] [输出.xlsx]
默认：python 批量跟进消息生成.py 示例客户数据.xlsx 跟进消息输出.xlsx
依赖：openpyxl（pip install openpyxl）
对应：轻流「AI 应用解决方案实习生」JD 加分项——自动化脚本 / Python / Excel 数据处理
说明：本脚本用规则模板生成（零 AI 依赖、可批量稳定跑）；如需更自然的文案，
      可把 build_message() 替换为调用大模型 API（传入同样字段即可）。
"""
import sys
from pathlib import Path
from openpyxl import load_workbook, Workbook

REQUIRED_COLS = ["客户名称", "公司", "行业", "岗位", "意向等级", "客户留言"]


def build_message(row: dict) -> str:
    """按意向等级选择话术模板，结合行业/留言做个性化微调"""
    name = (row.get("客户名称") or "客户").strip()
    company = (row.get("公司") or "贵司").strip()
    industry = (row.get("行业") or "").strip()
    note = (row.get("客户留言") or "").strip()
    grade = (row.get("意向等级") or "").strip().upper()

    if grade == "SQL":
        hook = f"您提到的「{note[:18]}…」需求我们已整理成方案要点" if len(note) > 18 else (
            "您提到的需求我们已整理成方案要点" if note else "您关注的能力点我们已有成熟方案")
        return (f"{name}您好，感谢关注！针对{company}的情况，{hook}。"
                f"我们可以先安排 20 分钟在线演示，结合{(industry or '同行业')}的落地案例讲清楚实施路径，"
                f"明天上午或下午哪个时间方便？")

    if grade == "MQL":
        return (f"{name}您好，感谢关注～给您准备了一份{(industry or '同行业')}客户的 AI CRM 落地案例，"
                f"3 分钟就能看完，先帮您建立整体认知。后续有具体问题随时回复这条消息，"
                f"我帮您对接方案顾问，不打扰您工作。")

    return (f"{name}您好，已为您送上产品资料包，方便您随时查阅。"
            f"如果后续有正式需求，欢迎随时联系我们做针对性演示。")


def read_rows(path: str) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("输入 Excel 为空")
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    missing = [c for c in REQUIRED_COLS if c not in headers]
    if missing:
        raise SystemExit(f"缺少必需列：{missing}，实际表头：{headers}")
    out = []
    for r in rows[1:]:
        if all(c is None for c in r):
            continue
        out.append({h: r[i] for i, h in enumerate(headers) if i < len(r)})
    return out


def write_output(rows: list[dict], messages: list[str], path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "跟进消息"
    ws.append(REQUIRED_COLS + ["生成的跟进消息"])
    for row, msg in zip(rows, messages):
        ws.append([row.get(c) for c in REQUIRED_COLS] + [msg])
    # 列宽自适应（近似）
    widths = [14, 18, 12, 18, 10, 40, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    wb.save(path)


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "示例客户数据.xlsx"
    dst = sys.argv[2] if len(sys.argv) > 2 else "跟进消息输出.xlsx"
    if not Path(src).exists():
        raise SystemExit(f"找不到输入文件：{src}")

    rows = read_rows(src)
    messages = [build_message(r) for r in rows]
    write_output(rows, messages, dst)

    print(f"已处理 {len(rows)} 条客户 -> {dst}")
    print("-" * 72)
    for r, m in zip(rows, messages):
        grade = (r.get("意向等级") or "?").strip()
        print(f"[{grade}] {r.get('客户名称')} · {r.get('公司')}")
        print(f"    {m[:46]}…")
    print("-" * 72)
    print("完成。输出文件新增「生成的跟进消息」列。")


if __name__ == "__main__":
    main()
